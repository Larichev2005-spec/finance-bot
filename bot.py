"""
Финансовый помощник для Telegram.
Быстрый ввод:  "кофе 250"  |  "такси 480"  |  "+50000 зарплата"  |  "1.5к продукты"
Минус = расход (по умолчанию), плюс в начале = доход.
Запуск: python bot.py  (токен в переменной BOT_TOKEN или в файле .env)
"""
import asyncio
import io
import os
import re
import sqlite3
from datetime import datetime

from aiogram import Bot, Dispatcher, F
from aiogram.filters import Command
from aiogram.types import (BufferedInputFile, CallbackQuery, InlineKeyboardButton,
                           InlineKeyboardMarkup, KeyboardButton, Message,
                           ReplyKeyboardMarkup)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------- Настройки ----------
if os.path.exists(".env"):
    for line in open(".env", encoding="utf-8"):
        if "=" in line and not line.startswith("#"):
            k, v = line.strip().split("=", 1)
            os.environ.setdefault(k, v)

TOKEN = os.getenv("BOT_TOKEN", "")
CUR = os.getenv("CURRENCY", "₽")
DB = os.path.join(os.getenv("DATA_DIR", "."), "finance.db")


# Цвета таблицы Excel: тёмно-синий, белый, чёрный
NAVY, WHITE, BLACK, LIGHT = "0B1F4B", "FFFFFF", "000000", "EEF1F7"

EXPENSE_CATS = {
    "🛒 Продукты": ["продукт", "магнит", "пятерочка", "перекресток", "лента", "ашан", "вкусвилл", "хлеб", "молоко"],
    "☕ Кафе и еда": ["кофе", "кафе", "обед", "ужин", "завтрак", "пицца", "ресторан", "бургер", "суши", "доставка"],
    "🚕 Транспорт": ["такси", "метро", "автобус", "бензин", "проезд", "парковка", "яндекс go"],
    "🏠 Дом и счета": ["аренда", "квартира", "жкх", "свет", "вода", "интернет", "связь", "телефон"],
    "💊 Здоровье": ["аптека", "врач", "лекарств", "стоматолог", "анализ"],
    "🎮 Досуг": ["кино", "игра", "подписка", "netflix", "spotify", "концерт", "бар", "клуб"],
    "👕 Покупки": ["одежда", "обувь", "wb", "ozon", "озон", "вайлдберриз", "магазин", "техника"],
    "📦 Прочее": [],
}
INCOME_CATS = {
    "💼 Зарплата": ["зарплата", "зп", "аванс", "оклад", "премия"],
    "🧾 Подработка": ["подработка", "фриланс", "заказ", "проект"],
    "🎁 Подарок": ["подарок", "подарили"],
    "💰 Прочий доход": [],
}

KB = ReplyKeyboardMarkup(
    keyboard=[
        [KeyboardButton(text="📊 Сводка"), KeyboardButton(text="📋 Последние")],
        [KeyboardButton(text="📁 Таблица Excel"), KeyboardButton(text="↩️ Отменить")],
    ],
    resize_keyboard=True,
    input_field_placeholder="кофе 250  или  +50000 зарплата",
)

NUM = re.compile(r"(?<![\w.,])(\d+(?:[.,]\d+)?)\s*([кk])?(?![\w])", re.I)


# ---------- База данных ----------
def db():
    con = sqlite3.connect(DB)
    con.row_factory = sqlite3.Row
    return con


def init_db():
    with db() as con:
        con.execute(
            """CREATE TABLE IF NOT EXISTS ops (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                user_id INTEGER NOT NULL,
                kind TEXT NOT NULL,           -- 'expense' | 'income'
                amount REAL NOT NULL,
                category TEXT NOT NULL,
                note TEXT,
                created TEXT NOT NULL
            )"""
        )


# ---------- Разбор сообщения ----------
def guess_category(text: str, kind: str) -> str:
    cats = INCOME_CATS if kind == "income" else EXPENSE_CATS
    low = text.lower()
    for cat, words in cats.items():
        if any(w in low for w in words):
            return cat
    return list(cats)[-1]


def parse(text: str):
    text = text.strip()
    m = NUM.search(text)
    if not m:
        return None
    amount = float(m.group(1).replace(",", "."))
    if m.group(2):
        amount *= 1000
    if amount <= 0:
        return None
    kind = "income" if text.startswith("+") else "expense"
    note = (text[: m.start()] + " " + text[m.end():]).replace("+", " ").strip(" -–—")
    note = re.sub(r"\s+", " ", note)
    return kind, amount, note, guess_category(note, kind)


def money(x: float) -> str:
    return f"{x:,.0f}".replace(",", " ") + f" {CUR}"


# ---------- Бот ----------
dp = Dispatcher()


def op_kb(op_id: int) -> InlineKeyboardMarkup:
    return InlineKeyboardMarkup(inline_keyboard=[[
        InlineKeyboardButton(text="🏷 Категория", callback_data=f"cat:{op_id}"),
        InlineKeyboardButton(text="↩️ Отменить", callback_data=f"undo:{op_id}"),
    ]])


@dp.message(Command("start", "help"))
async def start(m: Message):
    await m.answer(
        "<b>Финансовый помощник</b>\n\n"
        "Просто пиши операции одной строкой:\n"
        "▫️ <code>кофе 250</code> — расход\n"
        "▫️ <code>такси 480</code> — расход\n"
        "▫️ <code>+50000 зарплата</code> — доход\n"
        "▫️ <code>1.5к продукты</code> — 1 500\n\n"
        "Категория определяется сама, её можно поменять кнопкой.\n"
        "Всё хранится в таблице — выгрузить её можно в один тап.",
        parse_mode="HTML", reply_markup=KB,
    )


@dp.message(F.text == "📊 Сводка")
@dp.message(Command("stats"))
async def stats(m: Message):
    month = datetime.now().strftime("%Y-%m")
    with db() as con:
        rows = con.execute(
            "SELECT kind, category, SUM(amount) s FROM ops "
            "WHERE user_id=? AND created LIKE ? GROUP BY kind, category ORDER BY s DESC",
            (m.from_user.id, month + "%"),
        ).fetchall()
    inc = sum(r["s"] for r in rows if r["kind"] == "income")
    exp = sum(r["s"] for r in rows if r["kind"] == "expense")
    lines = [
        f"<b>{datetime.now():%m.%Y}</b>",
        f"⚪ Доходы: <b>{money(inc)}</b>",
        f"⚫ Расходы: <b>{money(exp)}</b>",
        f"🔵 Баланс: <b>{money(inc - exp)}</b>",
    ]
    exp_rows = [r for r in rows if r["kind"] == "expense"]
    if exp_rows:
        lines.append("\n<b>Куда уходят деньги</b>")
        for r in exp_rows[:6]:
            share = r["s"] / exp if exp else 0
            bar = "▰" * round(share * 10) + "▱" * (10 - round(share * 10))
            lines.append(f"{r['category']}\n<code>{bar}</code> {money(r['s'])} · {share:.0%}")
    await m.answer("\n".join(lines), parse_mode="HTML")


@dp.message(F.text == "📋 Последние")
async def last(m: Message):
    with db() as con:
        rows = con.execute(
            "SELECT * FROM ops WHERE user_id=? ORDER BY id DESC LIMIT 10", (m.from_user.id,)
        ).fetchall()
    if not rows:
        return await m.answer("Пока пусто. Напиши, например: <code>кофе 250</code>", parse_mode="HTML")
    out = []
    for r in rows:
        sign = "+" if r["kind"] == "income" else "−"
        d = datetime.fromisoformat(r["created"]).strftime("%d.%m")
        out.append(f"{d} · {r['category']} · <b>{sign}{money(r['amount'])}</b>"
                   + (f"\n      <i>{r['note']}</i>" if r["note"] else ""))
    await m.answer("\n".join(out), parse_mode="HTML")


@dp.message(F.text == "↩️ Отменить")
async def undo_last(m: Message):
    with db() as con:
        r = con.execute("SELECT id, amount FROM ops WHERE user_id=? ORDER BY id DESC LIMIT 1",
                        (m.from_user.id,)).fetchone()
        if not r:
            return await m.answer("Отменять нечего.")
        con.execute("DELETE FROM ops WHERE id=?", (r["id"],))
    await m.answer(f"Удалила последнюю запись на {money(r['amount'])}.")


@dp.message(F.text == "📁 Таблица Excel")
@dp.message(Command("export"))
async def export(m: Message):
    with db() as con:
        ops = con.execute("SELECT * FROM ops WHERE user_id=? ORDER BY created", (m.from_user.id,)).fetchall()
    if not ops:
        return await m.answer("Таблица пока пустая.")

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor=NAVY)
    head_font = Font(bold=True, color=WHITE, name="Calibri", size=11)
    band = PatternFill("solid", fgColor=LIGHT)

    def style_head(ws, n):
        for c in range(1, n + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "Операции"
    ws.append(["Дата", "Тип", "Категория", "Сумма", "Комментарий"])
    style_head(ws, 5)
    for i, r in enumerate(ops, start=2):
        income = r["kind"] == "income"
        ws.append([datetime.fromisoformat(r["created"]), "Доход" if income else "Расход",
                   r["category"], r["amount"] if income else -r["amount"], r["note"] or ""])
        ws.cell(row=i, column=1).number_format = "DD.MM.YYYY HH:MM"
        ws.cell(row=i, column=4).number_format = f'#,##0 "{CUR}"'
        ws.cell(row=i, column=4).font = Font(bold=True, color=NAVY if income else BLACK)
        if i % 2 == 1:
            for c in range(1, 6):
                ws.cell(row=i, column=c).fill = band
    for col, w in zip("ABCDE", (18, 10, 22, 16, 36)):
        ws.column_dimensions[col].width = w

    # Сводка по месяцам
    ws2 = wb.create_sheet("По месяцам")
    ws2.append(["Месяц", "Доходы", "Расходы", "Баланс"])
    style_head(ws2, 4)
    months = {}
    for r in ops:
        k = r["created"][:7]
        d = months.setdefault(k, [0, 0])
        d[0 if r["kind"] == "income" else 1] += r["amount"]
    for i, (k, (inc, exp)) in enumerate(sorted(months.items()), start=2):
        ws2.append([k, inc, exp, f"=B{i}-C{i}"])
        for c in (2, 3, 4):
            ws2.cell(row=i, column=c).number_format = f'#,##0 "{CUR}"'
        ws2.cell(row=i, column=4).font = Font(bold=True, color=NAVY)
    for c in range(1, 5):
        ws2.column_dimensions[get_column_letter(c)].width = 16

    # Сводка по категориям расходов
    ws3 = wb.create_sheet("Категории")
    ws3.append(["Категория", "Расходы"])
    style_head(ws3, 2)
    cats = {}
    for r in ops:
        if r["kind"] == "expense":
            cats[r["category"]] = cats.get(r["category"], 0) + r["amount"]
    for i, (k, v) in enumerate(sorted(cats.items(), key=lambda x: -x[1]), start=2):
        ws3.append([k, v])
        ws3.cell(row=i, column=2).number_format = f'#,##0 "{CUR}"'
    ws3.column_dimensions["A"].width = 24
    ws3.column_dimensions["B"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    name = f"finance_{datetime.now():%Y-%m-%d}.xlsx"
    await m.answer_document(BufferedInputFile(buf.getvalue(), filename=name), caption="Твоя таблица 📁")


@dp.callback_query(F.data.startswith("undo:"))
async def cb_undo(c: CallbackQuery):
    op_id = int(c.data.split(":")[1])
    with db() as con:
        con.execute("DELETE FROM ops WHERE id=? AND user_id=?", (op_id, c.from_user.id))
    await c.message.edit_text("↩️ Запись удалена.")
    await c.answer()


@dp.callback_query(F.data.startswith("cat:"))
async def cb_cat(c: CallbackQuery):
    op_id = int(c.data.split(":")[1])
    with db() as con:
        r = con.execute("SELECT kind FROM ops WHERE id=? AND user_id=?", (op_id, c.from_user.id)).fetchone()
    if not r:
        return await c.answer("Запись не найдена")
    cats = list(INCOME_CATS if r["kind"] == "income" else EXPENSE_CATS)
    rows = [[InlineKeyboardButton(text=cat, callback_data=f"set:{op_id}:{i}")] for i, cat in enumerate(cats)]
    await c.message.edit_reply_markup(reply_markup=InlineKeyboardMarkup(inline_keyboard=rows))
    await c.answer()


@dp.callback_query(F.data.startswith("set:"))
async def cb_set(c: CallbackQuery):
    _, op_id, idx = c.data.split(":")
    with db() as con:
        r = con.execute("SELECT * FROM ops WHERE id=? AND user_id=?", (op_id, c.from_user.id)).fetchone()
        if not r:
            return await c.answer("Запись не найдена")
        cats = list(INCOME_CATS if r["kind"] == "income" else EXPENSE_CATS)
        cat = cats[int(idx)]
        con.execute("UPDATE ops SET category=? WHERE id=?", (cat, op_id))
    await c.message.edit_text(card(r["kind"], r["amount"], cat, r["note"]), parse_mode="HTML",
                              reply_markup=op_kb(int(op_id)))
    await c.answer("Готово")


def card(kind, amount, cat, note):
    icon, sign = ("⚪", "+") if kind == "income" else ("⚫", "−")
    txt = f"{icon} <b>{sign}{money(amount)}</b>\n{cat}"
    return txt + (f"\n<i>{note}</i>" if note else "")


@dp.message(F.text)
async def quick_add(m: Message):
    p = parse(m.text)
    if not p:
        return await m.answer("Не нашла сумму 🤔\nНапиши так: <code>кофе 250</code> или <code>+50000 зарплата</code>",
                              parse_mode="HTML")
    kind, amount, note, cat = p
    with db() as con:
        cur = con.execute(
            "INSERT INTO ops (user_id, kind, amount, category, note, created) VALUES (?,?,?,?,?,?)",
            (m.from_user.id, kind, amount, cat, note, datetime.now().isoformat(timespec="seconds")),
        )
        op_id = cur.lastrowid
    await m.answer(card(kind, amount, cat, note), parse_mode="HTML", reply_markup=op_kb(op_id))


async def main():
    if not TOKEN:
        raise SystemExit("Укажи BOT_TOKEN (получить у @BotFather)")
    init_db()
    await dp.start_polling(Bot(TOKEN))


if __name__ == "__main__":
    asyncio.run(main())
